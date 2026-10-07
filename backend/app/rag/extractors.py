import io
import re
from typing import List, Dict, Any

def clean_text(text: str) -> str:
    """Normalize whitespace and strip unprintable characters."""
    if not text:
        return ""
    # Replace null bytes and strange controls
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    # Replace multiple spaces/tabs
    text = re.sub(r'[ \t]+', ' ', text)
    # Replace more than two consecutive newlines with double newline
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_text_from_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract page-wise text from PDF using pypdf."""
    pages = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            cleaned = clean_text(text)
            if cleaned:
                pages.append({
                    "page_number": idx + 1,
                    "section": f"Page {idx + 1}",
                    "text": cleaned
                })
    except Exception as e:
        # Fallback raw extraction if pypdf encounters an issue
        text = file_bytes.decode("utf-8", errors="ignore")
        cleaned = clean_text(text)
        if cleaned:
            pages.append({"page_number": 1, "section": "Document", "text": cleaned})
    return pages

def extract_text_from_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract section/paragraph text from DOCX."""
    pages = []
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_paras = []
        current_section = "General"
        page_num = 1

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and 'Heading' in p.style.name:
                current_section = text
            full_paras.append(text)
            if len(full_paras) >= 15:
                chunk_text = clean_text("\n\n".join(full_paras))
                if chunk_text:
                    pages.append({
                        "page_number": page_num,
                        "section": current_section,
                        "text": chunk_text
                    })
                    page_num += 1
                full_paras = []

        if full_paras:
            chunk_text = clean_text("\n\n".join(full_paras))
            if chunk_text:
                pages.append({
                    "page_number": page_num,
                    "section": current_section,
                    "text": chunk_text
                })
    except Exception as e:
        text = file_bytes.decode("utf-8", errors="ignore")
        cleaned = clean_text(text)
        if cleaned:
            pages.append({"page_number": 1, "section": "Document", "text": cleaned})
    return pages

def extract_text_from_pptx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract slide-wise text from PPTX."""
    pages = []
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            title = f"Slide {idx + 1}"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                if len(slide_texts) > 0 and len(slide_texts[0]) < 80:
                    title = slide_texts[0]
                cleaned = clean_text("\n".join(slide_texts))
                if cleaned:
                    pages.append({
                        "page_number": idx + 1,
                        "section": title,
                        "text": cleaned
                    })
    except Exception as e:
        text = file_bytes.decode("utf-8", errors="ignore")
        cleaned = clean_text(text)
        if cleaned:
            pages.append({"page_number": 1, "section": "Presentation", "text": cleaned})
    return pages

def extract_text_from_txt(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extract text from plain text file."""
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="ignore")

    cleaned = clean_text(text)
    if not cleaned:
        return []

    # Split into logical sections/pages if long
    paragraphs = [p.strip() for p in cleaned.split("\n\n") if p.strip()]
    pages = []
    current_page = []
    page_num = 1
    char_count = 0

    for p in paragraphs:
        current_page.append(p)
        char_count += len(p)
        if char_count >= 1500:
            pages.append({
                "page_number": page_num,
                "section": f"Section {page_num}",
                "text": "\n\n".join(current_page)
            })
            page_num += 1
            current_page = []
            char_count = 0

    if current_page:
        pages.append({
            "page_number": page_num,
            "section": f"Section {page_num}",
            "text": "\n\n".join(current_page)
        })

    return pages

def extract_document_text(filename: str, file_bytes: bytes) -> List[Dict[str, Any]]:
    """Router to extract text by file extension."""
    ext = filename.lower().split('.')[-1]
    if ext == 'pdf':
        return extract_text_from_pdf(file_bytes)
    elif ext in ['docx', 'doc']:
        return extract_text_from_docx(file_bytes)
    elif ext in ['pptx', 'ppt']:
        return extract_text_from_pptx(file_bytes)
    elif ext in ['txt', 'md']:
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: .{ext}. Supported formats are: PDF, DOCX, PPTX, TXT")
